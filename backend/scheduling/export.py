"""Shareable exports (Excel / PDF) of the current schedule and the installation plan."""
import re
from collections import defaultdict
from copy import deepcopy
from datetime import date as CalendarDate, timedelta
from io import BytesIO
from typing import Literal
from xml.sax.saxutils import escape

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

DAYS = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
MONTHS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
SCHEDULE_HEAD = ["Fecha", "Día", "Molde", "Panel", "Tipo", "Fachada", "Planta", "Frente", "Semana", "m²"]
PLAN_HEAD = ["Orden", "Frente", "Semana", "Panel", "Molde", "Tipo", "Fachada", "Fecha programada", "m²"]


def fmt(value):
    return CalendarDate.fromisoformat(value).strftime("%d/%m/%Y") if value else "Sin fecha"


def build_tables(data):
    """Summary, schedule rows (by date) and installation plan rows (by selection order)."""
    panels = data["panels"]
    scheduled = sorted((p for p in panels if p.get("date")), key=lambda p: (p["date"], p["molde"], p["code"]))
    schedule_rows = [[fmt(p["date"]), DAYS[CalendarDate.fromisoformat(p["date"]).weekday()], p["molde"], p["code"],
                      p.get("tipo") or "", p.get("facade_label") or "", p.get("floor_label") or "",
                      p.get("front") or "", fmt(p["week"]) if p.get("week") else "", p.get("area") or 0]
                     for p in scheduled]
    plan = sorted((p for p in panels if p.get("front")), key=lambda p: p["phase_order"])
    plan_rows = [[i + 1, p["front"], fmt(p["week"]), p["code"], p["molde"], p.get("tipo") or "",
                  p.get("facade_label") or "", fmt(p.get("date")), p.get("area") or 0]
                 for i, p in enumerate(plan)]
    area = round(sum(p.get("area") or 0 for p in scheduled), 2)
    summary = [["Inicio", fmt(data.get("start_date"))], ["Primera fecha", fmt(data.get("first_date"))],
               ["Fecha final", fmt(data.get("finish_date"))], ["Objetivo diario", f"{data.get('daily_capacity')} paneles/día"],
               ["Paneles programados", data.get("scheduled", 0)], ["Paneles sin fecha", data.get("unscheduled", 0)],
               ["Paneles en plan de instalación", len(plan_rows)], ["m² programados", area]]
    return summary, schedule_rows, plan_rows


def to_xlsx(summary, schedule_rows, plan_rows):
    wb = Workbook()
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1C1C1E")
    sheets = [("Resumen", ["Dato", "Valor"], summary), ("Cronograma", SCHEDULE_HEAD, schedule_rows),
              ("Plan de instalación", PLAN_HEAD, plan_rows)]
    for index, (title, head, rows) in enumerate(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        ws.append(head)
        for cell in ws[1]:
            cell.font, cell.fill = bold, fill
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(max(len(str(c.value or "")) for c in column) + 2, 40)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def pdf_table(head, rows):
    rows = rows or [["—"] + [""] * (len(head) - 1)]
    table = Table([head] + [[str(v) for v in r] for r in rows], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1C1C1E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F7")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C7C7CC")),
    ]))
    return table


def to_pdf(summary, schedule_rows, plan_rows):
    out = BytesIO()
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title="Cronograma y plan de instalación")
    doc.build([
        Paragraph("Cronograma y plan de instalación", styles["Title"]),
        Paragraph(f"Generado el {CalendarDate.today().strftime('%d/%m/%Y')}", styles["Normal"]), Spacer(1, 10),
        pdf_table(["Dato", "Valor"], summary), PageBreak(),
        Paragraph("Cronograma", styles["Heading2"]), pdf_table(SCHEDULE_HEAD, schedule_rows), PageBreak(),
        Paragraph("Plan de instalación", styles["Heading2"]), pdf_table(PLAN_HEAD, plan_rows),
    ])
    return out.getvalue()


GANTT_DAYS = 14


def tint(hex_color, alpha=0.2):
    """Mold color blended toward white, like the on-screen Gantt cells."""
    base = colors.HexColor(hex_color) if isinstance(hex_color, str) and len(hex_color) == 7 else colors.HexColor("#F2F2F7")
    return colors.Color(1 - (1 - base.red) * alpha, 1 - (1 - base.green) * alpha, 1 - (1 - base.blue) * alpha)


def to_gantt_pdf(data, single=False):
    """Mold rows × day columns: one sheet per block of GANTT_DAYS days, or every day on a single sheet."""
    panels = [p for p in data["panels"] if p.get("date")]
    cells = defaultdict(list)
    for panel in panels:
        cells[(panel["molde"], panel["date"])].append(panel)
    natural = lambda name: [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", name)]
    molds = sorted(data["molds"], key=lambda m: natural(m["name"]))
    styles = getSampleStyleSheet()
    out = BytesIO()
    title = Paragraph("Diagrama de Gantt · Cronograma de fabricación", styles["Title"])
    if not panels:
        SimpleDocTemplate(out, pagesize=landscape(A4), title="Diagrama de Gantt").build(
            [title, Paragraph("No hay paneles programados.", styles["Normal"])])
        return out.getvalue()
    start, finish = CalendarDate.fromisoformat(data["first_date"]), CalendarDate.fromisoformat(data["finish_date"])
    span = (finish - start).days + 1 if single else GANTT_DAYS
    label_width = 3.2 * cm
    # Multi-sheet keeps A4 width; a single sheet widens the page so every day keeps a readable column.
    day_width = (landscape(A4)[0] - 2 * cm - 12 - label_width) / GANTT_DAYS
    day_width = min(day_width, (14400 - 2 * cm - 12 - label_width) / span)  # PDF viewers cap pages at 200 in.
    page_width = max(landscape(A4)[0], label_width + day_width * span + 2 * cm + 12)
    pages = []
    block = start
    while block <= finish:
        days = [block + timedelta(days=i) for i in range(span)]
        block += timedelta(days=span)
        table = gantt_table(days, molds, cells, data["daily_capacity"], styles, label_width, day_width)
        if table:
            pages.append([Paragraph(f"{days[0].strftime('%d/%m/%Y')} – {days[-1].strftime('%d/%m/%Y')}", styles["Heading3"]), table])
    # Fit each complete block on its own sheet, even when molds have many copies.
    pages[0].insert(0, title)
    available_width = page_width - 2 * cm - 12  # ReportLab frame padding.
    content_height = max(sum(deepcopy(part).wrap(available_width, 100000)[1] + part.getSpaceBefore() + part.getSpaceAfter()
                             for part in page) for page in pages)
    page_height = max(landscape(A4)[1], content_height + 2 * cm + 48)
    doc = SimpleDocTemplate(out, pagesize=(page_width, page_height), leftMargin=cm, rightMargin=cm,
                            topMargin=cm, bottomMargin=cm, title="Diagrama de Gantt")
    story = []
    for page in pages:
        if story:
            story.append(PageBreak())
        story.extend(page)
    doc.build(story)
    return out.getvalue()


def gantt_table(days, molds, cells, capacity, styles, label_width, day_width):
    """Table for one range of days, or None when nothing is produced in it."""
    keys = [d.isoformat() for d in days]
    if not any((m["name"], k) in cells for m in molds for k in keys):
        return None  # skip blocks without production
    cell_style = styles["BodyText"].clone("gantt", fontSize=6.5, leading=8, alignment=0)
    mold_style = styles["BodyText"].clone("gantt-mold", fontSize=8, leading=10)
    day_style = styles["BodyText"].clone("gantt-day", fontSize=8, leading=10, alignment=1)
    day_style.textColor = colors.HexColor("#1C1C1E")
    head = [Paragraph("<b>Molde</b><br/><font size='7' color='#636366'>Programados / total</font>", mold_style)] + [
        Paragraph(f"<font size='7' color='#636366'>{DAYS[d.weekday()]}</font><br/>"
                  f"<b><font size='12'>{d.day}</font></b><br/>"
                  f"<font size='7' color='#636366'>{MONTHS[d.month - 1]}</font>", day_style)
        for d in days]
    rows, style = [head], [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FAFAFA")),
        ("BACKGROUND", (0, 1), (0, -1), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (1, 1), (-1, -1), 3),
        ("RIGHTPADDING", (1, 1), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E5EA")),
    ]
    for mold in molds:
        for copy in range(mold.get("copies", 1)):
            r = len(rows)
            color = mold.get("color") if re.fullmatch(r"#[0-9a-fA-F]{6}", mold.get("color") or "") else "#8E8E93"
            label = escape(mold["name"]) + (f" · copia {copy + 1}" if copy else "")
            details = (f"<br/><font size='7' color='#636366'>{mold.get('scheduled', 0)} / {mold.get('total', 0)} · "
                       f"{escape(mold.get('tipo') or 'Sin tipo')}</font>") if not copy else ""
            row = [Paragraph(f"<b>{label}</b>{details}", mold_style)]
            style.append(("LINEBEFORE", (0, r), (0, r), 3, colors.HexColor(color)))
            for c, key in enumerate(keys, start=1):
                group = cells.get((mold["name"], key), [])
                p = group[copy] if copy < len(group) else None
                if p:
                    code = p["code"]
                    quantity = re.search(r"\[([^\]]*)\]", code)
                    area = f"{p.get('area') or 0:g}".replace(".", ",")
                    row.append(Paragraph(f"<b>{escape(code.split(' [', 1)[0])}</b><br/>"
                                         f"<font size='7' color='#636366'>{escape(quantity.group(1) if quantity else '1 panel')}</font><br/>"
                                         f"<font size='7' color='#007AFF'><b>{area} m²</b></font>", cell_style))
                    panel_color = p.get("color") if re.fullmatch(r"#[0-9a-fA-F]{6}", p.get("color") or "") else color
                    style += [("BACKGROUND", (c, r), (c, r), tint(panel_color, 0.13)),
                              ("LINEBEFORE", (c, r), (c, r), 2, colors.HexColor(panel_color))]
                else:
                    row.append("")
                    if days[c - 1].weekday() == 6:
                        style.append(("BACKGROUND", (c, r), (c, r), colors.HexColor("#F2F2F7")))
            rows.append(row)
    counts = [sum(len(cells[(m["name"], k)]) for m in molds) for k in keys]
    areas = [round(sum(p.get("area") or 0 for m in molds for p in cells[(m["name"], k)]), 2) for k in keys]
    rows.append(["Paneles / día"] + [f"{n} / {0 if days[i].weekday() == 6 else capacity}" for i, n in enumerate(counts)])
    rows.append(["m² / día"] + [f"{a:g}".replace(".", ",") + " m²" if a else "—" for a in areas])
    style += [("BACKGROUND", (0, -2), (-1, -1), colors.HexColor("#F2F2F7")),
              ("TEXTCOLOR", (1, -1), (-1, -1), colors.HexColor("#007AFF")),
              ("FONTNAME", (0, -2), (-1, -1), "Helvetica-Bold"),
              ("LINEABOVE", (0, -2), (-1, -2), 1, colors.HexColor("#C7C7CC"))]
    table = Table(rows, colWidths=[label_width] + [day_width] * len(days), repeatRows=1)
    table.setStyle(TableStyle(style))
    return table


FORMATS = {"xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
           "pdf": (to_pdf, "application/pdf")}


def create_export_router(service):
    router = APIRouter(prefix="/schedule", tags=["Cronograma"])

    @router.get("/export.{kind}")
    async def export_schedule(kind: str):
        if kind not in FORMATS:
            raise HTTPException(404, "Formato no soportado")
        data = await service.get()
        data = data.model_dump() if hasattr(data, "model_dump") else data
        build, media = FORMATS[kind]
        name = f"cronograma_{CalendarDate.today().isoformat()}.{kind}"
        return StreamingResponse(BytesIO(build(*build_tables(data))), media_type=media,
                                 headers={"Content-Disposition": f'attachment; filename="{name}"'})

    @router.get("/export-gantt.pdf")
    async def export_gantt(layout: Literal["pages", "single"] = "pages"):
        data = await service.get()
        data = data.model_dump() if hasattr(data, "model_dump") else data
        name = f"gantt_{CalendarDate.today().isoformat()}.pdf"
        return StreamingResponse(BytesIO(to_gantt_pdf(data, single=layout == "single")), media_type="application/pdf",
                                 headers={"Content-Disposition": f'attachment; filename="{name}"'})

    return router
