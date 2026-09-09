"""Regression tests for reports dashboard APIs and export flows."""
import os
from io import BytesIO

import pytest
import requests
from openpyxl import load_workbook


def _base_url() -> str:
    env_base = os.environ.get("REACT_APP_BACKEND_URL")
    if env_base:
        return env_base.rstrip("/")
    with open("/app/frontend/.env", "r", encoding="utf-8") as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL="):
                return line.split("=", 1)[1].strip().rstrip("/")
    raise RuntimeError("REACT_APP_BACKEND_URL is required for tests")


BASE_URL = _base_url()
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def api_client():
    """Shared HTTP session for report endpoints."""
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def admin_token(api_client):
    """Admin auth token fixture for protected write endpoints."""
    res = api_client.post(f"{API}/admin/verify", json={"password": "admin2026"}, timeout=30)
    assert res.status_code == 200
    body = res.json()
    assert body.get("ok") is True
    token = body.get("token")
    assert isinstance(token, str) and token
    return token


def test_report_molds_baseline_shape_and_consistency(api_client):
    """Reports module: baseline /report/molds data and self-consistent counters."""
    res = api_client.get(f"{API}/report/molds", timeout=30)
    assert res.status_code == 200
    data = res.json()

    assert isinstance(data["total"], int)
    assert isinstance(data["con_molde"], int)
    assert isinstance(data["sin_molde"], int)
    assert data["total"] == data["con_molde"] + data["sin_molde"]
    assert data["total"] == len(data["items"])

    resumen_sum = sum(r["count"] for r in data["resumen"])
    assert resumen_sum == data["con_molde"]

    for row in data["resumen"]:
        assert {"molde", "tipo", "color", "count"}.issubset(row.keys())
        assert isinstance(row["count"], int)


@pytest.mark.parametrize("facade", ["all", "norte", "sur", "este", "oeste"])
def test_report_molds_facade_filter_supported_values(api_client, facade):
    """Reports filters: supported facade keys in lowercase."""
    res = api_client.get(f"{API}/report/molds", params={"facade": facade}, timeout=30)
    assert res.status_code == 200
    data = res.json()
    assert data["total"] == len(data["items"])
    if facade != "all":
        assert all(i["facade"] == facade for i in data["items"])


def test_report_molds_unassigned_semantics(api_client):
    """Reports filters: molde=__none__ returns only unassigned panels."""
    res = api_client.get(f"{API}/report/molds", params={"molde": "__none__"}, timeout=30)
    assert res.status_code == 200
    data = res.json()
    assert data["con_molde"] == 0
    assert data["sin_molde"] == data["total"]
    assert data["resumen"] == []
    assert all(item["molde"] is None for item in data["items"])


def test_report_molds_tipo_none_semantics(api_client):
    """Reports filters: tipo=__none__ must exclude typed molds."""
    res = api_client.get(f"{API}/report/molds", params={"tipo": "__none__"}, timeout=30)
    assert res.status_code == 200
    data = res.json()
    assert data["total"] == len(data["items"])
    assert all(item["tipo"] is None for item in data["items"])


def test_report_molds_unmatched_compound_filter_returns_empty(api_client):
    """Reports filters: unmatched molde+tipo combination returns zero rows."""
    all_molds = api_client.get(f"{API}/molds", timeout=30)
    assert all_molds.status_code == 200
    items = all_molds.json().get("items", [])
    assert items

    picked = items[0]
    wrong_tipo = "__QA_NOT_A_REAL_TYPE__"
    res = api_client.get(
        f"{API}/report/molds",
        params={"molde": picked["name"], "tipo": wrong_tipo},
        timeout=30,
    )
    assert res.status_code == 200
    data = res.json()
    assert data["total"] == 0
    assert data["items"] == []
    assert data["resumen"] == []


@pytest.mark.parametrize(
    "fmt,magic,content_type",
    [
        ("pdf", b"%PDF", "application/pdf"),
        ("xlsx", b"PK", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ],
)
def test_report_export_files_by_format(api_client, fmt, magic, content_type):
    """Export module: PDF/XLSX bytes and response metadata."""
    res = api_client.get(
        f"{API}/report/molds/export",
        params={"format": fmt, "facade": "norte", "molde": "all", "tipo": "all"},
        timeout=60,
    )
    assert res.status_code == 200
    assert content_type in res.headers.get("content-type", "")
    assert "attachment" in res.headers.get("content-disposition", "").lower()
    assert res.content[: len(magic)] == magic


def test_report_export_xlsx_contains_kpis(api_client):
    """Export module: XLSX has KPI line and data rows for selected filters."""
    report = api_client.get(f"{API}/report/molds", params={"facade": "all"}, timeout=30).json()
    res = api_client.get(
        f"{API}/report/molds/export",
        params={"format": "xlsx", "facade": "all", "molde": "all", "tipo": "all"},
        timeout=60,
    )
    assert res.status_code == 200

    wb = load_workbook(BytesIO(res.content), read_only=True)
    ws = wb.active
    kpi_line = [cell for cell in next(ws.iter_rows(min_row=3, max_row=3, values_only=True)) if cell]
    assert any(str(report["total"]) in str(v) for v in kpi_line)
    assert any(str(report["con_molde"]) in str(v) for v in kpi_line)
    assert any(str(report["sin_molde"]) in str(v) for v in kpi_line)


def test_report_export_xlsx_filtered_row_count_matches_api(api_client):
    """Export module: filtered XLSX carries exactly filtered API item rows."""
    params = {"facade": "este", "molde": "__none__", "tipo": "all"}
    report = api_client.get(f"{API}/report/molds", params=params, timeout=30)
    assert report.status_code == 200
    report_data = report.json()

    export_res = api_client.get(
        f"{API}/report/molds/export",
        params={"format": "xlsx", **params},
        timeout=60,
    )
    assert export_res.status_code == 200
    wb = load_workbook(BytesIO(export_res.content), read_only=True)
    ws = wb.active

    header_row = None
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        if row and row[0] == "Pieza":
            header_row = idx
            break
    assert header_row is not None

    item_rows = 0
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row and any(cell not in (None, "") for cell in row):
            item_rows += 1
    assert item_rows == report_data["total"]


def test_write_endpoint_requires_admin_token(api_client, admin_token):
    """Auth guard: protected write endpoint denies missing token and accepts valid admin token."""
    unauth = api_client.post(f"{API}/tipos", json={"name": "QA_TMP_AUTH_CHECK"}, timeout=30)
    assert unauth.status_code == 401

    auth = api_client.post(
        f"{API}/tipos",
        json={"name": "QA_TMP_AUTH_CHECK"},
        headers={"Authorization": f"Bearer {admin_token}"},
        timeout=30,
    )
    assert auth.status_code == 200
    assert auth.json()["name"] == "QA_TMP_AUTH_CHECK"

    cleanup = api_client.delete(
        f"{API}/tipos/QA_TMP_AUTH_CHECK",
        headers={"Authorization": f"Bearer {admin_token}"},
        timeout=30,
    )
    assert cleanup.status_code == 200
    assert cleanup.json()["deleted"] is True
