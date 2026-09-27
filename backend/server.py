from fastapi import FastAPI, APIRouter, HTTPException, Query, UploadFile, File
from fastapi.responses import FileResponse, Response
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import json
import re
import uuid
import logging
import requests
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict, BeforeValidator
from typing import List, Optional, Annotated
from datetime import datetime, timezone
from project_panels import ProjectPanelsService, ProjectPanelsResponse, ProjectPanelsUpdate
from scheduling.service import ScheduleService
from scheduling.routes import create_schedule_router
from scheduling.export import create_export_router
from scheduling.installation_export import create_installation_export_router
from scheduling.report_exports import create_report_exports_router
from scheduling.spatial import SpatialCatalog
from scheduling.phases import PhaseStore, create_phase_router
from backup import create_backup_router
from auth import create_auth_router, require_admin_for_writes

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

STATIC_DIR = ROOT_DIR / 'static'
MODEL_PATH = STATIC_DIR / 'nab3d.glb'
OBJECTS_CACHE = STATIC_DIR / 'objects.json'
VIEWER_PATH = STATIC_DIR / 'viewer.html'

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

# ---------- Object storage (Emergent) ----------

STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
APP_NAME = "bimtracker"
storage_key = None


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": os.environ.get("EMERGENT_LLM_KEY")}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}",
                        headers={"X-Storage-Key": key, "Content-Type": content_type},
                        data=data, timeout=120)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(f"{STORAGE_URL}/objects/{path}",
                            headers={"X-Storage-Key": key, "Content-Type": content_type},
                            data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()


def storage_get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")

# ---------- Object catalog (extracted from GLB) ----------

def load_objects():
    if OBJECTS_CACHE.exists():
        with open(OBJECTS_CACHE) as f:
            return json.load(f)
    from pygltflib import GLTF2
    g = GLTF2().load(str(MODEL_PATH))
    objs, seen = [], set()
    for n in g.nodes:
        if n.mesh is not None and n.name and n.name not in seen:
            seen.add(n.name)
            objs.append({'name': n.name, 'mark': n.name.split(' ')[0]})
    with open(OBJECTS_CACHE, 'w') as f:
        json.dump(objs, f)
    return objs


import mmap
import struct
import numpy as np

GLB: dict = {}
MESH_CACHE: dict = {}
COMP_DTYPE = {5120: np.int8, 5121: np.uint8, 5122: np.int16, 5123: np.uint16, 5125: np.uint32, 5126: np.float32}
TYPE_SIZE = {"SCALAR": 1, "VEC2": 2, "VEC3": 3, "VEC4": 4, "MAT4": 16}


def load_glb_index():
    """Parse GLB header + JSON chunk once; keep the binary chunk memory-mapped."""
    f = open(MODEL_PATH, 'rb')
    mm = mmap.mmap(f.fileno(), 0, access=mmap.ACCESS_READ)
    clen, _ = struct.unpack_from('<II', mm, 12)
    js = json.loads(mm[20:20 + clen])
    bin_off = 20 + clen + 8
    parents = {}
    for i, n in enumerate(js.get('nodes', [])):
        for c in n.get('children', []):
            parents[c] = i
    by_name = {}
    for i, n in enumerate(js.get('nodes', [])):
        if n.get('mesh') is not None and n.get('name') and n['name'] not in by_name:
            by_name[n['name']] = i
    GLB.update(js=js, mm=mm, bin_off=bin_off, parents=parents, by_name=by_name)


def node_local_matrix(n: dict) -> np.ndarray:
    if n.get('matrix'):
        return np.array(n['matrix'], dtype=np.float64).reshape(4, 4).T
    t = n.get('translation', [0, 0, 0])
    x, y, z, w = n.get('rotation', [0, 0, 0, 1])
    s = n.get('scale', [1, 1, 1])
    R = np.array([
        [1 - 2 * (y * y + z * z), 2 * (x * y - z * w), 2 * (x * z + y * w)],
        [2 * (x * y + z * w), 1 - 2 * (x * x + z * z), 2 * (y * z - x * w)],
        [2 * (x * z - y * w), 2 * (y * z + x * w), 1 - 2 * (x * x + y * y)],
    ])
    M = np.eye(4)
    M[:3, :3] = R * np.array(s)[None, :]
    M[:3, 3] = t
    return M


def node_world_matrix(idx: int) -> np.ndarray:
    M = node_local_matrix(GLB['js']['nodes'][idx])
    p = GLB['parents'].get(idx)
    while p is not None:
        M = node_local_matrix(GLB['js']['nodes'][p]) @ M
        p = GLB['parents'].get(p)
    return M


def read_accessor(idx: int) -> np.ndarray:
    js = GLB['js']
    acc = js['accessors'][idx]
    bv = js['bufferViews'][acc['bufferView']]
    dtype = np.dtype(COMP_DTYPE[acc['componentType']])
    ncomp = TYPE_SIZE[acc['type']]
    count = acc['count']
    start = GLB['bin_off'] + bv.get('byteOffset', 0) + acc.get('byteOffset', 0)
    stride = bv.get('byteStride') or dtype.itemsize * ncomp
    if stride == dtype.itemsize * ncomp:
        arr = np.frombuffer(GLB['mm'], dtype=dtype, count=count * ncomp, offset=start)
    else:
        raw = np.frombuffer(GLB['mm'], dtype=np.uint8, count=stride * count, offset=start)
        arr = np.lib.stride_tricks.as_strided(raw.view(dtype), shape=(count, ncomp), strides=(stride, dtype.itemsize)).copy()
    return arr.reshape(count, ncomp)


def extract_mesh(name: str) -> dict:
    if name in MESH_CACHE:
        return MESH_CACHE[name]
    idx = GLB['by_name'].get(name)
    if idx is None:
        raise HTTPException(status_code=404, detail="Geometría no encontrada")
    js = GLB['js']
    M = node_world_matrix(idx)
    positions, indices, base = [], [], 0
    for prim in js['meshes'][js['nodes'][idx]['mesh']]['primitives']:
        if prim.get('mode', 4) != 4 or 'POSITION' not in prim['attributes']:
            continue
        pos = read_accessor(prim['attributes']['POSITION']).astype(np.float64)
        pos = pos @ M[:3, :3].T + M[:3, 3]
        if prim.get('indices') is not None:
            ind = read_accessor(prim['indices']).reshape(-1).astype(np.int64)
        else:
            ind = np.arange(len(pos), dtype=np.int64)
        positions.append(pos)
        indices.append(ind + base)
        base += len(pos)
    if not positions:
        raise HTTPException(status_code=404, detail="Geometría vacía")
    P = np.vstack(positions)
    center = (P.min(axis=0) + P.max(axis=0)) / 2
    P = P - center
    result = {
        "name": name,
        "positions": [round(float(v), 4) for v in P.reshape(-1)],
        "indices": np.concatenate(indices).tolist(),
        "size": [round(float(v), 4) for v in (P.max(axis=0) - P.min(axis=0))],
    }
    if len(MESH_CACHE) > 200:
        MESH_CACHE.clear()
    MESH_CACHE[name] = result
    return result


OBJECTS: List[dict] = []
NAME_SET = set()
FACADE_NAMES = set()
FACADES: dict = {}
FACADES_PATH = STATIC_DIR / 'facades.json'
DIMS: dict = {}
DIMS_PATH = STATIC_DIR / 'dims.json'
VALID_FACADES = {"norte", "sur", "este", "oeste"}
FACADE_LABELS = {"norte": "Norte", "sur": "Sur", "este": "Este", "oeste": "Oeste", "all": "Todas"}

# Façade = marks starting with C or L followed by a number (C1, C7B-2, L3G, ...) → 533 objects (per user)
FACADE_RE = re.compile(r"^[CL]\d")


def is_facade(mark: str) -> bool:
    return bool(FACADE_RE.match(mark))


@app.on_event("startup")
async def startup():
    global OBJECTS, NAME_SET, FACADE_NAMES, FACADES, DIMS
    OBJECTS = load_objects()
    NAME_SET = {o['name'] for o in OBJECTS}
    try:
        load_glb_index()
    except Exception as e:
        logging.error(f"GLB index failed: {e}")
    FACADE_NAMES = {o['name'] for o in OBJECTS if is_facade(o['mark'])}
    if FACADES_PATH.exists():
        with open(FACADES_PATH) as f:
            FACADES = json.load(f)
    if DIMS_PATH.exists():
        with open(DIMS_PATH) as f:
            DIMS = json.load(f)
    await db.tags.create_index("object_name", unique=True)
    await db.molds.create_index("name", unique=True)
    await db.tipos.create_index("name", unique=True)
    # One-time migration: the old status-based tagging model was replaced by molds.
    # Legacy tag docs (status/observations) carry no meaningful data for the new model.
    await db.tags.delete_many({"status": {"$exists": True}})
    # Medidas/color moved from the tag to the mold catalog: drop the obsolete per-tag fields.
    await db.tags.update_many({}, {"$unset": {"ancho": "", "alto": "", "color": ""}})
    if await db.tipos.count_documents({}) == 0:
        await db.tipos.insert_many([{"name": n} for n in ["Curvo", "Liso", "Borde de losa", "Cubre viga"]])
    if not os.environ.get("EMERGENT_LLM_KEY"):
        logging.warning("Object storage not configured (EMERGENT_LLM_KEY missing); photo/PDF uploads are disabled")
    else:
        try:
            init_storage()
            logging.info("Object storage initialized")
        except Exception as e:
            logging.warning(f"Storage init failed: {e}")
    logging.info(f"Loaded {len(OBJECTS)} objects ({len(FACADE_NAMES)} facade) from model")


# ---------- Models ----------

PyObjectId = Annotated[str, BeforeValidator(str)]


class BaseDocument(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    id: Optional[PyObjectId] = Field(default=None, alias="_id")

    def to_mongo(self) -> dict:
        d = self.model_dump(by_alias=True)
        d.pop("_id", None)
        return d

    @classmethod
    def from_mongo(cls, doc: Optional[dict]):
        return cls(**doc) if doc else None


class Tag(BaseDocument):
    object_name: str
    molde: Optional[str] = None
    notas: str = ""
    photo: Optional[str] = None
    updated_at: str = ""
    created_at: str = ""
    history: List[dict] = []


class TagUpsert(BaseModel):
    object_name: str
    molde: Optional[str] = None
    notas: str = ""
    photo: Optional[str] = None


class BulkTagUpsert(BaseModel):
    object_names: List[str]
    molde: Optional[str] = None
    notas: str = ""
    photo: Optional[str] = None


class Mold(BaseDocument):
    name: str
    tipo: Optional[str] = None
    color: str
    ancho: Optional[float] = None
    alto: Optional[float] = None
    photo: Optional[str] = None
    plano_pdf: Optional[str] = None


class MoldUpsert(BaseModel):
    name: str
    tipo: str
    color: str
    ancho: Optional[float] = None
    alto: Optional[float] = None
    photo: Optional[str] = None
    plano_pdf: Optional[str] = None


class Tipo(BaseDocument):
    name: str


class TipoUpsert(BaseModel):
    name: str


class FacadesPayload(BaseModel):
    facades: dict


class DimsPayload(BaseModel):
    dims: dict


# ---------- Helpers ----------

def display_name(name: str) -> str:
    parts = name.split(" ")
    if len(parts) >= 2 and parts[0] == parts[1]:
        parts.pop(1)
    return " ".join(parts)


async def fetch_tags_map() -> dict:
    tags = {}
    async for doc in db.tags.find():
        try:
            t = Tag.from_mongo(doc)
        except Exception:
            logging.warning(f"Skipping malformed tag doc: {doc.get('object_name')}")
            continue
        tags[t.object_name] = t
    return tags


async def fetch_molds_map() -> dict:
    molds = {}
    async for doc in db.molds.find():
        m = Mold.from_mongo(doc)
        molds[m.name] = m
    return molds


# ---------- Routes ----------

project_panels = ProjectPanelsService(db, lambda: FACADE_NAMES, fetch_tags_map)

@api_router.get("/")
async def root():
    return {"message": "NaBiMoldes API", "objects": len(OBJECTS)}


@api_router.get("/model")
async def get_model():
    return FileResponse(MODEL_PATH, media_type="model/gltf-binary",
                        headers={"Cache-Control": "public, max-age=86400"})


@api_router.get("/viewer")
async def get_viewer():
    return FileResponse(VIEWER_PATH, media_type="text/html",
                        headers={"Cache-Control": "no-cache"})


@api_router.get("/facades/count")
async def facades_count():
    return {"count": len(FACADES)}


@api_router.get("/project/panels", response_model=ProjectPanelsResponse)
async def get_project_panels():
    return await project_panels.read()


@api_router.put("/project/panels", response_model=ProjectPanelsResponse)
async def update_project_panels(payload: ProjectPanelsUpdate):
    return await project_panels.save(payload)


@api_router.post("/facades")
async def save_facades(payload: FacadesPayload):
    """Persist the per-object cardinal orientation computed by the 3D viewer."""
    global FACADES
    clean = {k: v for k, v in payload.facades.items() if k in NAME_SET and v in VALID_FACADES}
    if not clean:
        raise HTTPException(status_code=422, detail="Sin orientaciones válidas")
    FACADES = clean
    with open(FACADES_PATH, 'w') as f:
        json.dump(clean, f)
    return {"saved": len(clean)}


@api_router.get("/dims/count")
async def dims_count():
    return {"count": len(DIMS)}


@api_router.post("/dims")
async def save_dims(payload: DimsPayload):
    """Persist per-object bounding box sizes [sx, sy, sz] computed by the 3D viewer."""
    global DIMS
    clean = {}
    for k, v in payload.dims.items():
        if k in NAME_SET and isinstance(v, list) and len(v) == 3:
            try:
                clean[k] = [round(float(x), 4) for x in v]
            except (TypeError, ValueError):
                continue
    if not clean:
        raise HTTPException(status_code=422, detail="Sin dimensiones válidas")
    DIMS = clean
    with open(DIMS_PATH, 'w') as f:
        json.dump(clean, f)
    return {"saved": len(clean)}


@api_router.get("/tipos")
async def list_tipos():
    items = [doc["name"] async for doc in db.tipos.find().sort("name", 1)]
    return {"items": items}


@api_router.post("/tipos")
async def create_tipo(payload: TipoUpsert):
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=422, detail="El nombre del tipo es obligatorio")
    existing = await db.tipos.find_one({"name": name})
    if not existing:
        await db.tipos.insert_one({"name": name})
    return {"name": name}


@api_router.put("/tipos/{name}")
async def rename_tipo(name: str, payload: TipoUpsert):
    new_name = payload.name.strip()
    if not new_name:
        raise HTTPException(status_code=422, detail="El nombre del tipo es obligatorio")
    if not await db.tipos.find_one({"name": name}):
        raise HTTPException(status_code=404, detail="Tipo no encontrado")
    await db.tipos.update_one({"name": name}, {"$set": {"name": new_name}})
    await db.molds.update_many({"tipo": name}, {"$set": {"tipo": new_name}})
    return {"name": new_name}


@api_router.delete("/tipos/{name}")
async def delete_tipo(name: str):
    await db.tipos.delete_one({"name": name})
    await db.molds.update_many({"tipo": name}, {"$set": {"tipo": None}})
    return {"deleted": True, "name": name}


@api_router.get("/molds")
async def list_molds():
    molds = await fetch_molds_map()
    return {"items": [
        {"name": m.name, "tipo": m.tipo, "color": m.color, "ancho": m.ancho, "alto": m.alto,
         "photo": m.photo, "plano_pdf": m.plano_pdf}
        for m in molds.values()
    ]}


@api_router.post("/molds")
async def save_mold(payload: MoldUpsert):
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=422, detail="El nombre del molde es obligatorio")
    tipo = (payload.tipo or "").strip()
    if not tipo:
        raise HTTPException(status_code=422, detail="El tipo de molde es obligatorio")
    # Reuse an existing tipo case-insensitively; otherwise auto-register the new one.
    existing = await db.tipos.find_one({"name": {"$regex": f"^{re.escape(tipo)}$", "$options": "i"}})
    if existing:
        tipo = existing["name"]
    else:
        await db.tipos.insert_one({"name": tipo})
    payload.tipo = tipo
    color = (payload.color or "").strip() or "#8E8E93"
    mold = Mold(name=name, tipo=payload.tipo, color=color, ancho=payload.ancho, alto=payload.alto,
                photo=payload.photo, plano_pdf=payload.plano_pdf)
    await db.molds.update_one({"name": name}, {"$set": mold.to_mongo()}, upsert=True)
    return {"name": mold.name, "tipo": mold.tipo, "color": mold.color, "ancho": mold.ancho,
            "alto": mold.alto, "photo": mold.photo, "plano_pdf": mold.plano_pdf}


@api_router.delete("/molds/{name}")
async def delete_mold(name: str):
    await db.molds.delete_one({"name": name})
    now = datetime.now(timezone.utc).isoformat()
    await db.tags.update_many(
        {"molde": name},
        {"$set": {"molde": None, "updated_at": now}},
    )
    return {"deleted": True, "name": name}


ALLOWED_IMG = {"image/jpeg", "image/png", "image/webp", "image/gif", "image/heic", "image/heif"}


@api_router.post("/upload")
async def upload_photo(file: UploadFile = File(...)):
    if file.content_type not in ALLOWED_IMG:
        raise HTTPException(status_code=422, detail="Solo se permiten imágenes (JPG, PNG, WEBP, GIF)")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="La imagen supera el límite de 10 MB")
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in (file.filename or "") else "jpg"
    path = f"{APP_NAME}/uploads/{uuid.uuid4()}.{ext}"
    try:
        result = put_object(path, data, file.content_type)
    except Exception as e:
        logging.error(f"Photo upload failed: {e}")
        raise HTTPException(status_code=502, detail="No se pudo subir la foto. Inténtalo de nuevo.")
    await db.files.insert_one({
        "storage_path": result["path"],
        "original_filename": file.filename,
        "content_type": file.content_type,
        "size": result.get("size", len(data)),
        "is_deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    return {"path": result["path"]}


@api_router.post("/upload/pdf")
async def upload_pdf(file: UploadFile = File(...)):
    if file.content_type != "application/pdf":
        raise HTTPException(status_code=422, detail="Solo se permiten archivos PDF")
    data = await file.read()
    if len(data) > 25 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="El PDF supera el límite de 25 MB")
    path = f"{APP_NAME}/planos/{uuid.uuid4()}.pdf"
    try:
        result = put_object(path, data, "application/pdf")
    except Exception as e:
        logging.error(f"PDF upload failed: {e}")
        raise HTTPException(status_code=502, detail="No se pudo subir el PDF. Inténtalo de nuevo.")
    await db.files.insert_one({
        "storage_path": result["path"],
        "original_filename": file.filename,
        "content_type": "application/pdf",
        "size": result.get("size", len(data)),
        "is_deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    return {"path": result["path"], "filename": file.filename}


@api_router.get("/files/{path:path}")
async def serve_file(path: str):
    record = await db.files.find_one({"storage_path": path, "is_deleted": False})
    if not record:
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    try:
        data, ct = storage_get_object(path)
    except Exception as e:
        logging.error(f"Photo download failed: {e}")
        raise HTTPException(status_code=502, detail="No se pudo descargar la foto.")
    return Response(content=data, media_type=record.get("content_type") or ct,
                    headers={"Cache-Control": "public, max-age=86400"})


@api_router.get("/photos")
async def list_photos(
    facade: str = "all",
    from_: str = Query(default="", alias="from"),
    to: str = "",
):
    """Gallery of all obra photos attached to tags, filterable by facade and date."""
    if facade != "all" and facade not in VALID_FACADES:
        raise HTTPException(status_code=422, detail="Fachada inválida")
    molds = await fetch_molds_map()
    items = []
    async for doc in db.tags.find({"photo": {"$ne": None}}):
        t = Tag.from_mongo(doc)
        if not t.photo:
            continue
        fac = FACADES.get(t.object_name)
        if facade != "all" and fac != facade:
            continue
        d = (t.updated_at or "")[:10]
        if from_ and d < from_:
            continue
        if to and d > to:
            continue
        mold = molds.get(t.molde) if t.molde else None
        items.append({
            "name": t.object_name,
            "mark": t.object_name.split(" ")[0],
            "facade": fac,
            "photo": t.photo,
            "text": t.notas,
            "date": t.updated_at,
            "molde": t.molde,
            "tipo": mold.tipo if mold else None,
            "color_molde": mold.color if mold else None,
        })
    items.sort(key=lambda x: x["date"] or "", reverse=True)
    return {"total": len(items), "items": items}


@api_router.delete("/photos")
async def delete_photo(object_name: str, photo: str):
    """Remove the obra photo from a tag and soft-delete its file record."""
    doc = await db.tags.find_one({"object_name": object_name})
    t = Tag.from_mongo(doc)
    if not t or t.photo != photo:
        raise HTTPException(status_code=404, detail="Foto no encontrada")
    if not t.molde and not t.notas:
        await db.tags.delete_one({"object_name": object_name})
    else:
        await db.tags.update_one(
            {"object_name": object_name},
            {"$set": {"photo": None, "updated_at": datetime.now(timezone.utc).isoformat()}},
        )
    await db.files.update_one({"storage_path": photo}, {"$set": {"is_deleted": True}})
    return {"deleted": True, "object_name": object_name}


@api_router.get("/objects")
async def list_objects(
    search: str = "",
    molde: str = "all",
    facade: str = "all",
    skip: int = 0,
    limit: int = Query(default=50, le=200),
):
    tags = await fetch_tags_map()
    molds = await fetch_molds_map()
    q = search.strip().lower()
    filtered = []
    for o in OBJECTS:
        if q and q not in o['name'].lower():
            continue
        if facade != "all" and FACADES.get(o['name']) != facade:
            continue
        t = tags.get(o['name'])
        t_molde = t.molde if t else None
        if molde == "none":
            if t_molde is not None:
                continue
        elif molde != "all":
            if t_molde != molde:
                continue
        mold = molds.get(t_molde) if t_molde else None
        filtered.append({
            "name": o['name'],
            "mark": o['mark'],
            "facade": FACADES.get(o['name']),
            "molde": t_molde,
            "tipo": mold.tipo if mold else None,
            "color_molde": mold.color if mold else None,
            "notas": t.notas if t else "",
        })
    total = len(filtered)
    return {"total": total, "items": filtered[skip:skip + limit]}


@api_router.get("/objects/names")
async def list_object_names(facade: str = "all", molde: str = "all"):
    """Names of facade panels matching the filters (for bulk selection in the 3D viewer)."""
    if facade != "all" and facade not in VALID_FACADES:
        raise HTTPException(status_code=422, detail="Fachada no válida")
    tags = await fetch_tags_map()
    names = []
    for name in sorted(FACADE_NAMES):
        if facade != "all" and FACADES.get(name) != facade:
            continue
        t = tags.get(name)
        t_molde = t.molde if t else None
        if molde == "none":
            if t_molde is not None:
                continue
        elif molde != "all" and t_molde != molde:
            continue
        names.append(name)
    return {"names": names, "total": len(names)}


@api_router.get("/object")
async def get_object(name: str):
    if name not in NAME_SET:
        raise HTTPException(status_code=404, detail="Objeto no encontrado")
    doc = await db.tags.find_one({"object_name": name})
    t = Tag.from_mongo(doc)
    molds = await fetch_molds_map()
    mold = molds.get(t.molde) if t and t.molde else None
    mark = name.split(' ')[0]
    return {
        "name": name,
        "mark": mark,
        "facade": FACADES.get(name),
        "dimensions": DIMS.get(name),
        "molde": t.molde if t else None,
        "tipo": mold.tipo if mold else None,
        "color_molde": mold.color if mold else None,
        "ancho": mold.ancho if mold else None,
        "alto": mold.alto if mold else None,
        "notas": t.notas if t else "",
        "photo": t.photo if t else None,
        "created_at": t.created_at if t else "",
        "history": t.history if t else [],
    }


@api_router.get("/object/mesh")
async def get_object_mesh(name: str):
    """Triangulated geometry of a single object (world-space, centered) for the panel preview."""
    if name not in NAME_SET:
        raise HTTPException(status_code=404, detail="Objeto no encontrado")
    if not GLB:
        raise HTTPException(status_code=503, detail="Modelo no indexado")
    return extract_mesh(name)


@api_router.get("/tags")
async def get_tags():
    tags = await fetch_tags_map()
    molds = await fetch_molds_map()
    result = {}
    for name, t in tags.items():
        mold = molds.get(t.molde) if t.molde else None
        result[name] = {"molde": t.molde, "color_molde": mold.color if mold else None}
    return result


async def upsert_tag_doc(payload: TagUpsert) -> dict:
    if payload.object_name not in NAME_SET:
        raise HTTPException(status_code=404, detail="Objeto no encontrado")
    molde = (payload.molde or "").strip() or None
    notas = payload.notas.strip()
    photo = (payload.photo or "").strip() or None
    existing = await db.tags.find_one({"object_name": payload.object_name})
    prev = Tag.from_mongo(existing) if existing else None
    if not molde and not notas and not photo:
        await db.tags.delete_one({"object_name": payload.object_name})
        return {"object_name": payload.object_name, "molde": None,
                "notas": "", "photo": None, "history": prev.history if prev else []}
    now = datetime.now(timezone.utc).isoformat()
    created_at = prev.created_at if prev and prev.created_at else now
    history = list(prev.history) if prev else []
    changed = not prev or (prev.molde != molde or prev.notas != notas or prev.photo != photo)
    if changed:
        history.append({"molde": molde, "notas": notas, "photo": photo, "date": now})
    tag = Tag(
        object_name=payload.object_name,
        molde=molde,
        notas=notas,
        photo=photo,
        updated_at=now,
        created_at=created_at,
        history=history,
    )
    await db.tags.update_one(
        {"object_name": payload.object_name},
        {"$set": tag.to_mongo()},
        upsert=True,
    )
    return {
        "object_name": tag.object_name,
        "molde": tag.molde,
        "notas": tag.notas,
        "photo": tag.photo,
        "created_at": tag.created_at,
        "history": tag.history,
    }


@api_router.put("/tags")
async def upsert_tag(payload: TagUpsert):
    async with project_panels.write_lock:
        await project_panels.check_assignment([payload.object_name], payload.molde)
        return await upsert_tag_doc(payload)


@api_router.put("/tags/bulk")
async def bulk_upsert_tags(payload: BulkTagUpsert):
    names = list(dict.fromkeys(n for n in payload.object_names if n in NAME_SET))
    if not names:
        raise HTTPException(status_code=422, detail="Sin piezas válidas seleccionadas")
    updated = []
    async with project_panels.write_lock:
        await project_panels.check_assignment(names, payload.molde)
        for name in names:
            item = TagUpsert(object_name=name, molde=payload.molde, notas=payload.notas, photo=payload.photo)
            await upsert_tag_doc(item)
            updated.append(name)
    return {"updated": len(updated), "object_names": updated}


@api_router.delete("/tags")
async def delete_tag(object_name: str):
    if object_name not in NAME_SET:
        raise HTTPException(status_code=404, detail="Objeto no encontrado")
    doc = await db.tags.find_one({"object_name": object_name})
    if doc and doc.get("photo"):
        await db.files.update_one({"storage_path": doc["photo"]}, {"$set": {"is_deleted": True}})
    await db.tags.delete_one({"object_name": object_name})
    return {"deleted": True, "object_name": object_name}


async def build_molds_report(facade: str = "all", molde: str = "all", tipo: str = "all"):
    if facade != "all" and facade not in VALID_FACADES:
        raise HTTPException(status_code=422, detail="Fachada inválida")
    tags = await fetch_tags_map()
    molds = await fetch_molds_map()
    items = []
    resumen = {}
    con_molde = 0
    for name in FACADE_NAMES:
        fac = FACADES.get(name)
        if facade != "all" and fac != facade:
            continue
        t = tags.get(name)
        molde_name = t.molde if t else None
        mold = molds.get(molde_name) if molde_name else None
        tipo_name = mold.tipo if mold else None
        if molde == "__none__":
            if molde_name is not None:
                continue
        elif molde != "all" and molde_name != molde:
            continue
        if tipo == "__none__":
            if tipo_name is not None:
                continue
        elif tipo != "all" and tipo_name != tipo:
            continue
        if molde_name:
            con_molde += 1
            if molde_name not in resumen:
                resumen[molde_name] = {"molde": molde_name, "tipo": tipo_name,
                                       "color": mold.color if mold else None, "count": 0}
            resumen[molde_name]["count"] += 1
        items.append({
            "name": name,
            "mark": name.split(' ')[0],
            "facade": fac,
            "molde": molde_name,
            "tipo": tipo_name,
            "color": mold.color if mold else None,
            "ancho": mold.ancho if mold else None,
            "alto": mold.alto if mold else None,
        })
    items.sort(key=lambda x: (x["facade"] or "zzz", x["mark"], x["name"]))
    project = await project_panels.read(tags)
    scope = "project" if facade == molde == tipo == "all" else "filtered"
    total = project.total_panels if scope == "project" else len(items)
    return {
        "total": total,
        "con_molde": con_molde,
        "sin_molde": total - con_molde,
        "model_total": len(items),
        "model_unassigned": len(items) - con_molde,
        "scope": scope,
        "project": project.model_dump(),
        "items": items,
        "resumen": sorted(resumen.values(), key=lambda r: -r["count"]),
    }


@api_router.get("/report/molds")
async def get_molds_report(facade: str = "all", molde: str = "all", tipo: str = "all"):
    """Report of panels and their assigned mold, grouped/sorted by facade."""
    return await build_molds_report(facade, molde, tipo)


def _filter_label(facade: str, molde: str, tipo: str) -> str:
    parts = [f"Fachada: {FACADE_LABELS.get(facade, facade) if facade != 'all' else 'Todas'}"]
    if molde and molde != "all":
        parts.append(f"Molde: {'Sin molde' if molde == '__none__' else molde}")
    if tipo and tipo != "all":
        parts.append(f"Tipo: {'Sin tipo' if tipo == '__none__' else tipo}")
    return "   ·   ".join(parts)


def make_molds_pdf(data: dict, facade: str = "all", molde: str = "all", tipo: str = "all") -> bytes:
    from io import BytesIO
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors as rl_colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.platypus import Image as RLImage
    from reportlab.lib.styles import getSampleStyleSheet

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="Reporte de Moldes BIMTracker",
                            leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    elems = []
    logo_defs = [
        ("logo_fiberkret.png", 1032 / 290),
        ("logo_entrepisos.png", 1020 / 411),
        ("logo_grcontreras.png", 1340 / 542),
    ]
    lh = 10 * mm
    imgs = [RLImage(str(STATIC_DIR / f), width=lh * r, height=lh)
            for f, r in logo_defs if (STATIC_DIR / f).exists()]
    if imgs:
        letterhead = Table([imgs], colWidths=[i.drawWidth + 8 * mm for i in imgs], hAlign="CENTER")
        letterhead.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ]))
        elems.append(letterhead)
        elems.append(Spacer(1, 5 * mm))
    period = _filter_label(facade, molde, tipo)
    elems += [
        Paragraph("Reporte de Moldes — Paneles de Fachada", styles["Title"]),
        Paragraph(period, styles["Normal"]),
        Spacer(1, 4 * mm),
    ]
    pending_label = "Pendientes" if data["scope"] == "project" else "Sin molde"
    elems.append(Paragraph(
        f"Total paneles: {data['total']} &nbsp;·&nbsp; Con molde: {data['con_molde']} "
        f"&nbsp;·&nbsp; {pending_label}: {data['sin_molde']}", styles["Heading3"]))
    elems.append(Paragraph(
        f"Total del proyecto: {data['project']['total_panels']} "
        f"({'manual' if data['project']['is_manual'] else 'según modelo'}) "
        f"&nbsp;·&nbsp; Paneles del modelo en este reporte: {data['model_total']} "
        f"&nbsp;·&nbsp; Sin molde en el modelo: {data['model_unassigned']}", styles["Normal"]))
    if data["resumen"]:
        summary = "   ·   ".join(
            f"{r['molde']} ({r['tipo'] or '—'}): {r['count']}"
            for r in data["resumen"]
        )
        elems.append(Paragraph(summary, styles["Normal"]))
    elems.append(Spacer(1, 5 * mm))
    rows = [["Pieza", "Fachada", "Molde", "Tipo", "Medidas", "Color"]]
    color_cmds = []
    for i, it in enumerate(data["items"]):
        medidas = f"{it['ancho']} × {it['alto']}" if it.get("ancho") and it.get("alto") else "—"
        col = it.get("color")
        swatch = ""
        if col:
            try:
                color_cmds.append(("BACKGROUND", (5, i + 1), (5, i + 1), rl_colors.HexColor(col)))
            except Exception:
                swatch = "—"
        else:
            swatch = "—"
        rows.append([
            Paragraph(display_name(it["name"]), styles["BodyText"]),
            FACADE_LABELS.get(it.get("facade") or "", "—"),
            it["molde"] or "—",
            it["tipo"] or "—",
            medidas,
            swatch,
        ])
    table = Table(rows, colWidths=[46 * mm, 20 * mm, 28 * mm, 28 * mm, 24 * mm, 24 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#1C1C1E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, rl_colors.HexColor("#C7C7CC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [rl_colors.white, rl_colors.HexColor("#F2F2F7")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        *color_cmds,
    ]))
    elems.append(table)
    doc.build(elems)
    return buf.getvalue()


def make_molds_xlsx(data: dict, facade: str = "all", molde: str = "all", tipo: str = "all") -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Moldes"
    ws.append(["Reporte de Moldes — Paneles de Fachada"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([_filter_label(facade, molde, tipo)])
    ws.append([f"Total paneles: {data['total']}", f"Con molde: {data['con_molde']}", f"Sin molde: {data['sin_molde']}"])
    if data["scope"] == "project":
        ws.cell(row=3, column=3, value=f"Pendientes: {data['sin_molde']}")
    ws.append([f"Total del proyecto: {data['project']['total_panels']}",
               "Total manual" if data["project"]["is_manual"] else "Total según modelo",
               f"Paneles del modelo en este reporte: {data['model_total']}",
               f"Sin molde en el modelo: {data['model_unassigned']}"])
    if data["resumen"]:
        ws.append([f"{r['molde']} ({r['tipo'] or '—'})" for r in data["resumen"]])
        ws.append([str(r["count"]) for r in data["resumen"]])
    ws.append([])
    header = ["Pieza", "Fachada", "Molde", "Tipo", "Ancho", "Alto", "Color"]
    ws.append(header)
    hrow = ws.max_row
    for col in range(1, len(header) + 1):
        c = ws.cell(row=hrow, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1C1C1E", end_color="1C1C1E", fill_type="solid")
    for it in data["items"]:
        ws.append([
            display_name(it["name"]),
            FACADE_LABELS.get(it.get("facade") or "", "—"),
            it["molde"] or "—",
            it["tipo"] or "—",
            it.get("ancho") or "",
            it.get("alto") or "",
            it.get("color") or "",
        ])
    for col, width in zip("ABCDEFG", [40, 12, 20, 16, 10, 10, 16]):
        ws.column_dimensions[col].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@api_router.get("/report/molds/export")
async def export_molds_report(format: str = "xlsx", facade: str = "all", molde: str = "all", tipo: str = "all"):
    if format not in ("pdf", "xlsx"):
        raise HTTPException(status_code=422, detail="Formato inválido (pdf|xlsx)")
    data = await build_molds_report(facade, molde, tipo)
    if format == "pdf":
        content = make_molds_pdf(data, facade, molde, tipo)
        media = "application/pdf"
    else:
        content = make_molds_xlsx(data, facade, molde, tipo)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    filename = f"reporte_moldes_{facade}.{format}"
    return Response(
        content=content,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


schedule_spatial = SpatialCatalog(lambda: GLB, node_world_matrix, read_accessor)


AREA_CACHE: dict = {}


def mesh_surface_area(name: str) -> float:
    """Total surface of a panel's mesh patches (skin, ribs and cantos), in m²."""
    idx = GLB['by_name'].get(name)
    if idx is None:
        return 0.0
    M = node_world_matrix(idx)
    total = 0.0
    for prim in GLB['js']['meshes'][GLB['js']['nodes'][idx]['mesh']]['primitives']:
        if prim.get('mode', 4) != 4 or 'POSITION' not in prim['attributes']:
            continue
        pos = read_accessor(prim['attributes']['POSITION']).astype(np.float64) @ M[:3, :3].T + M[:3, 3]
        if prim.get('indices') is not None:
            ind = read_accessor(prim['indices']).reshape(-1).astype(np.int64)
        else:
            ind = np.arange(len(pos), dtype=np.int64)
        tris = pos[ind].reshape(-1, 3, 3)
        total += float(np.linalg.norm(np.cross(tris[:, 1] - tris[:, 0], tris[:, 2] - tris[:, 0]), axis=1).sum() / 2)
    return total


def panel_area(name):
    """Panel m² = half the total mesh surface. Approximation pending exact validation against
    Revit's solid surface area (~7% high); falls back to the bounding-box front face."""
    if name not in AREA_CACHE:
        area = round(mesh_surface_area(name) / 2, 2) if GLB else 0.0
        if not area:
            dims = DIMS.get(name)
            area = round(dims[1] * max(dims[0], dims[2]), 2) if dims else 0.0
        AREA_CACHE[name] = area
    return AREA_CACHE[name]


phase_store = PhaseStore(db)


async def schedule_source():
    phases = await phase_store.by_panel()
    tags = await fetch_tags_map()
    molds = await fetch_molds_map()
    project = await project_panels.read(tags)
    located, floors = schedule_spatial.locate(FACADE_NAMES, FACADES)
    panels = {}
    awaiting_location = 0
    eligible_count = 0
    for name in FACADE_NAMES:
        tag = tags.get(name)
        mold = molds.get(tag.molde) if tag and tag.molde else None
        if mold:
            eligible_count += 1
            if name not in located:
                awaiting_location += 1
                continue
            panels[name] = {"object_name": name, "code": display_name(name),
                            "molde": mold.name, "tipo": mold.tipo, "color": mold.color,
                            "area": panel_area(name), **phases.get(name, {}), **located[name]}
    return {"panels": panels, "total": project.total_panels,
            "floors": floors, "awaiting_location": awaiting_location, "eligible_count": eligible_count,
            "molds": [{"name": m.name, "tipo": m.tipo, "color": m.color} for m in molds.values()]}


schedule_service = ScheduleService(db, schedule_source)
api_router.include_router(create_export_router(schedule_service))
api_router.include_router(create_installation_export_router(phase_store))
api_router.include_router(create_report_exports_router())
api_router.include_router(create_schedule_router(schedule_service))
api_router.include_router(create_phase_router(phase_store, lambda: FACADE_NAMES))


def apply_restored_settings(settings):
    global FACADES, DIMS
    FACADES = settings["facades"]
    DIMS = settings["dims"]


api_router.include_router(create_backup_router(db, STATIC_DIR, storage_get_object, put_object, apply_restored_settings))
api_router.include_router(create_auth_router())
app.include_router(api_router)

# Registered before CORS so CORS stays outermost and 401 responses keep their CORS headers.
app.middleware("http")(require_admin_for_writes)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
